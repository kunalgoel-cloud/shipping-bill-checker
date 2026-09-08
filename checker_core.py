"""
checker_core.py
----------------
All non-UI logic for the Shipping Bill Checker:
  - reading the courier working sheet (auto header detection)
  - reading the commercial agreement workbook (zone matrix + zone map + charges)
  - deriving zones from state names
  - calculating the expected (agreement-side) price for an AWB
  - comparing it against the billed subtotal to decide Approved / Disputed
  - building the consolidated output workbook (with Disputed rows highlighted)

Kept separate from app.py so the logic can be unit-tested / reused without
Streamlit running.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from typing import Optional

import pandas as pd
from openpyxl.styles import PatternFill

# --------------------------------------------------------------------------
# Constants
# --------------------------------------------------------------------------

APPOINTMENT = "Appointment based Delivery"
NON_APPOINTMENT = "Non Appointment based Delivery"

STATUS_APPROVED = "Approved"
STATUS_DISPUTED = "Disputed"
STATUS_NOT_CHECKED = "Not Checked"

DISPUTED_FILL = "FFE0E0"  # light red, used both on-screen and in the Excel export

# Column names we ask the user to map from their working sheet.
#
# Required fields are everything the per-AWB price calculation and its
# billed-side comparison actually need to run — without any one of these the
# AWB can't be priced at all, so it's left "Not Checked" rather than guessed.
REQUIRED_FIELDS = {
    "awb": "AWB Number",
    "pickup_state": "Pickup State",
    "drop_state": "Drop State",
    "drop_location": "Delivery Location (City)",  # needed to detect metro drops
    "invoice_value": "Invoice Value",
    "freight_billed": "Freight Amount (Billed)",
    "fsc_billed": "Fuel Surcharge (Billed)",
    "fov_billed": "FOV (Billed)",
    "docket_billed": "Docket Charges (Billed)",
    "state_charges_billed": "State Charges (Billed)",  # billed counterpart of Metro Congestion Charge
    "appointment_billed": "Appointment Charges (Billed)",
}

# Optional: not needed to calculate/compare the price, but used for display,
# grouping, courier filtering, or as informational-only context in the "show
# the maths" breakdown (charges the agreement doesn't govern at all, so they
# never affect the Approved/Disputed decision).
OPTIONAL_FIELDS = {
    "courier": "Courier Name",
    "customer": "Customer / Client",
    "consignee": "Consignee Name",
    "chargeable_weight": "Chargeable Weight (Billed, reference only)",
    "box_count": "No of Boxes (Billed, reference only)",
    "rate_per_kg": "Rate per KG (Billed, reference only)",
    "total_billed": "Total Charges (Billed)",
    "dhp_billed": "DHP (Billed, informational)",
    "war_surcharge_billed": "War Surcharge (Billed, informational)",
    "oda_billed": "ODA Charges (Billed, informational)",
    "handling_billed": "Handling Charges (Billed, informational)",
    "green_tax_billed": "Green Tax (Billed, informational)",
    "demurrage_billed": "Demurrage Charge (Billed, informational)",
    "other_charges_billed": "Other Charges (Billed, informational)",
}

ALL_FIELDS = {**REQUIRED_FIELDS, **OPTIONAL_FIELDS}

# Fuzzy hints used to auto-suggest a column mapping from the uploaded sheet.
# Kept deliberately specific wherever a sheet might have a same-named
# "(Yes/No)" or otherwise unrelated column that could false-match a looser
# hint (e.g. "ODA Charges (Yes/No)" vs "ODA Charges (Billing)", or
# "Demurrage Days" vs "Demurrage Charge (Billing)").
AUTO_HINTS = {
    "awb": ["awb"],
    "pickup_state": ["pickup state", "origin state", "from state"],
    "drop_state": ["drop state", "destination state", "to state"],
    "drop_location": ["drop city", "destination city", "delivery city", "drop location"],
    "invoice_value": ["invoice value"],
    "freight_billed": ["freight amount (billing)", "freight amount", "freight (billing)"],
    "fsc_billed": ["fuel surcharge (billing)", "fuel surcharge", "fsc (billing)"],
    "fov_billed": ["fov (billing)", "fov"],
    "docket_billed": ["docket charges (billing)", "docket charge (billing)", "docket charges", "docket charge"],
    "state_charges_billed": ["state charges (billing)", "state charges"],
    "appointment_billed": ["appointment charges (billing)", "appointment charges billing"],
    "courier": ["courier"],
    "customer": ["client name", "customer"],
    # "consginee" covers the common courier-sheet misspelling of "consignee".
    "consignee": ["consignee name", "consginee name", "consignee", "consginee"],
    "chargeable_weight": ["chargeable weight"],
    "box_count": ["no of boxes", "no. of boxes", "number of boxes", "box count", "no of box"],
    "rate_per_kg": ["rate per kg", "rate/kg", "per kg rate"],
    "total_billed": ["total charges (billing)", "total charges"],
    "dhp_billed": ["dhp"],
    "war_surcharge_billed": ["war surcharge"],
    "oda_billed": ["oda charges (billing)", "oda (billing)"],
    "handling_billed": ["handling charges (billing)", "handling charges", "handling (billing)"],
    "green_tax_billed": ["green tax"],
    "demurrage_billed": ["demurrage charge (billing)", "demurrage charges (billing)"],
    "other_charges_billed": ["other charges"],
}


# --------------------------------------------------------------------------
# Working sheet loading
# --------------------------------------------------------------------------

def read_uploaded_table(uploaded_file) -> pd.DataFrame:
    """Read a csv/xlsx/xls upload with tolerant header detection.

    Many courier working sheets have a summary/total row above the real
    header, so we scan the first 10 rows for the one most likely to be a
    header (the row with the most non-empty text-looking cells that isn't
    mostly numeric).
    """
    name = (uploaded_file.name or "").lower()
    raw_bytes = uploaded_file.read()

    if name.endswith(".csv"):
        preview = pd.read_csv(io.BytesIO(raw_bytes), header=None, nrows=10, dtype=str)
    else:
        preview = pd.read_excel(io.BytesIO(raw_bytes), header=None, nrows=10, dtype=str)

    header_row = _guess_header_row(preview)

    if name.endswith(".csv"):
        df = pd.read_csv(io.BytesIO(raw_bytes), header=header_row)
    else:
        df = pd.read_excel(io.BytesIO(raw_bytes), header=header_row)

    # Drop fully-empty columns/rows and normalise column names to strings.
    df = df.dropna(axis=1, how="all").dropna(axis=0, how="all")
    df.columns = [str(c).strip() for c in df.columns]
    df = df.reset_index(drop=True)
    return df


def _guess_header_row(preview: pd.DataFrame) -> int:
    best_row, best_score = 0, -1
    for i in range(len(preview)):
        row = preview.iloc[i]
        non_null = row.notna().sum()
        text_like = sum(
            1 for v in row if isinstance(v, str) and not _looks_numeric(v)
        )
        score = non_null + text_like  # reward rows with lots of text labels
        if score > best_score:
            best_score = score
            best_row = i
    return best_row


def _looks_numeric(v: str) -> bool:
    v = v.strip().replace(",", "")
    if v == "":
        return False
    try:
        float(v)
        return True
    except ValueError:
        return False


def suggest_column_mapping(columns: list[str]) -> dict[str, Optional[str]]:
    """Best-effort auto mapping from sheet columns to our internal field names."""
    norm_cols = {c: re.sub(r"\s+", " ", c.strip().lower()) for c in columns}
    mapping: dict[str, Optional[str]] = {}
    for field_key, hints in AUTO_HINTS.items():
        found = None
        for col, norm in norm_cols.items():
            if any(h in norm for h in hints):
                found = col
                break
        mapping[field_key] = found
    return mapping


# --------------------------------------------------------------------------
# Commercial agreement loading
# --------------------------------------------------------------------------

@dataclass
class Agreement:
    courier_name: str
    zone_matrix: pd.DataFrame  # index = from-zone, columns = to-zone, values = rate/kg
    zone_map: dict[str, str]  # normalised state/location -> zone
    zone_locations: dict[str, list[str]] = field(default_factory=dict)  # zone -> original-cased locations, for display
    charges: dict[str, float] = field(default_factory=dict)
    metro_locations: list[str] = field(default_factory=list)
    source: str = "Uploaded"  # human-readable: where this agreement came from

    def rate_for(self, from_zone: str, to_zone: str) -> Optional[float]:
        try:
            return float(self.zone_matrix.loc[from_zone, to_zone])
        except (KeyError, ValueError, TypeError):
            return None

    def zone_for(self, state_or_city: Optional[str]) -> Optional[str]:
        if not state_or_city or pd.isna(state_or_city):
            return None
        key = _normalise(state_or_city)
        return self.zone_map.get(key)

    def is_metro(self, city: Optional[str]) -> bool:
        """Whether `city` counts as a metro location for congestion-charge
        purposes. Uses a contains-match (not exact) so real-world variants
        like "MUMBAI CITY" or "BANGALORE DELIVERY" still match "Mumbai" /
        "Bangalore" without needing the sheet to spell the city exactly as
        the agreement does.
        """
        if not city or pd.isna(city):
            return False
        key = _normalise(city)
        return any(_normalise(m) in key for m in self.metro_locations if m)


def _normalise(s: str) -> str:
    return re.sub(r"\s+", " ", str(s).strip().lower())


def _parse_zone_locations(rows: list[tuple[str, str]]) -> tuple[dict[str, str], dict[str, list[str]]]:
    """rows: (zone, "state1, state2, ...") pairs.

    Returns (normalised-location -> zone lookup map, zone -> original-cased
    locations list for display). Shared by load_agreement() and
    default_agreement() so both build identical structures.
    """
    zone_map: dict[str, str] = {}
    zone_locations: dict[str, list[str]] = {}
    for zone, locations in rows:
        zone = str(zone).strip()
        locations = str(locations)
        if not zone or locations.strip().lower() == "nan":
            continue
        locs = [loc.strip() for loc in locations.split(",") if loc.strip()]
        if not locs:
            continue
        zone_locations[zone] = locs
        for loc in locs:
            zone_map[_normalise(loc)] = zone
    return zone_map, zone_locations


def load_agreement(uploaded_file) -> Agreement:
    """Parse the agreement workbook (see build_template() for expected layout)."""
    raw = uploaded_file.read()
    xls = pd.ExcelFile(io.BytesIO(raw))

    matrix_df = pd.read_excel(xls, sheet_name="ZoneMatrix", index_col=0)
    matrix_df.index = [str(i).strip() for i in matrix_df.index]
    matrix_df.columns = [str(c).strip() for c in matrix_df.columns]

    mapping_df = pd.read_excel(xls, sheet_name="ZoneMapping")
    rows = [
        (row.get("Zone", ""), row.get("States / Locations (comma separated)", ""))
        for _, row in mapping_df.iterrows()
    ]
    zone_map, zone_locations = _parse_zone_locations(rows)

    charges_df = pd.read_excel(xls, sheet_name="Charges")
    charges: dict[str, float] = {}
    metro_locations: list[str] = []
    for _, row in charges_df.iterrows():
        param = str(row.get("Parameter", "")).strip()
        value = row.get("Value", "")
        if param == "Metro Locations (comma separated)":
            metro_locations = [x.strip() for x in str(value).split(",") if x.strip()]
        elif param:
            try:
                charges[param] = float(value)
            except (ValueError, TypeError):
                pass

    return Agreement(
        courier_name="Unknown",
        zone_matrix=matrix_df,
        zone_map=zone_map,
        zone_locations=zone_locations,
        charges=charges,
        metro_locations=metro_locations,
        source=f"Uploaded file: {getattr(uploaded_file, 'name', 'agreement.xlsx')}",
    )


# --------------------------------------------------------------------------
# Default (Safexpress) template — used both to seed the downloadable
# template and as a ready-to-go sample agreement.
# --------------------------------------------------------------------------

_ZONES = ["North-1", "North-2", "East", "North-East", "West-1", "West-2",
          "South-1", "South-2", "Central"]

_MATRIX_VALUES = [
    [6.48, 6.48, 10.8, 16.2, 7.56, 8.64, 8.64, 10.8, 7.56],
    [6.48, 6.48, 10.8, 16.2, 8.64, 8.64, 10.8, 10.8, 7.56],
    [8.64, 10.8, 6.48, 7.56, 8.64, 10.8, 8.64, 10.8, 7.56],
    [8.64, 10.8, 7.56, 6.48, 10.8, 10.8, 10.8, 16.2, 8.64],
    [7.56, 8.64, 10.8, 16.2, 6.48, 6.48, 8.64, 10.8, 7.56],
    [8.64, 10.8, 10.8, 16.2, 6.48, 6.48, 7.56, 10.8, 7.56],
    [8.64, 10.8, 10.8, 16.2, 8.64, 7.56, 6.48, 7.56, 7.56],
    [10.8, 10.8, 10.8, 16.2, 8.64, 8.64, 6.48, 6.48, 7.56],
    [7.56, 8.64, 10.8, 16.2, 6.48, 7.56, 7.56, 10.8, 6.48],
]

_ZONE_LOCATIONS = {
    "North-1": "Delhi, Uttar Pradesh, Haryana, Rajasthan",
    "North-2": "Chandigarh, Punjab, Himachal Pradesh, Uttarakhand, Jammu & Kashmir, Ladakh",
    "East": "West Bengal, Odisha, Bihar, Jharkhand, Chhattisgarh",
    "North-East": "Assam, Meghalaya, Tripura, Arunachal Pradesh, Mizoram, Manipur, Nagaland, Sikkim",
    "West-1": "Gujarat, Daman & Diu, Dadra & Nagar Haveli",
    "West-2": "Maharashtra, Goa",
    "South-1": "Andhra Pradesh, Telangana, Karnataka, Tamil Nadu, Puducherry",
    "South-2": "Kerala",
    "Central": "Madhya Pradesh",
}

# "Appointment Charge Amount (Rs)" is not in the source PDF — it's derived
# from the real working sheet, where every appointment-billed AWB carries
# exactly ₹750, consistently, so it's used as the concrete figure the price
# calculation adds when Appointment based Delivery is selected.
_DEFAULT_CHARGES = {
    "Docket Charge (Flat Rs)": 100,
    "FOV Percent (%)": 0.1,
    "FOV Minimum (Rs)": 100,
    "FSC Percent (%)": 20,
    "Min Chargeable Weight (Kg)": 15,
    "Min Chargeable Freight (Rs)": 400,
    "Metro Congestion Charge (Rs)": 100,
    "Appointment Charge Amount (Rs)": 750,
}

# "Bangalore" is included alongside "Bengaluru" because real working sheets
# commonly spell the city that way (e.g. "BANGALORE DELIVERY") — same city,
# different common spelling, not a judgment call.
_METRO_LOCATIONS = "Ahmedabad, Bengaluru, Bangalore, Chennai, Delhi, Hyderabad, Kolkata, Mumbai, Pune"


def default_agreement() -> Agreement:
    """The Safexpress commercial agreement, embedded directly in code (the
    constants above) rather than requiring a file upload.

    This is the agreement the app is active with from the moment it loads —
    no setup needed. Uploading a different agreement workbook in the
    Commercial Agreement tab overrides it for the session; nothing here is
    ever written back to disk, so the built-in figures below are the single
    source of truth and stay in sync with the downloadable template
    (build_template) and the Commercial Agreement view tab automatically.
    """
    matrix_df = pd.DataFrame(_MATRIX_VALUES, index=_ZONES, columns=_ZONES)
    matrix_df.index.name = "From \\ To"

    rows = [(z, _ZONE_LOCATIONS[z]) for z in _ZONES]
    zone_map, zone_locations = _parse_zone_locations(rows)

    metro_locations = [x.strip() for x in _METRO_LOCATIONS.split(",") if x.strip()]

    return Agreement(
        courier_name="Safexpress",
        zone_matrix=matrix_df,
        zone_map=zone_map,
        zone_locations=zone_locations,
        charges=dict(_DEFAULT_CHARGES),
        metro_locations=metro_locations,
        source="Built-in default (bundled in code) — Safexpress commercial agreement",
    )


def build_template(prefill: bool = True) -> bytes:
    """Build the commercial-agreement Excel template as bytes.

    If prefill=True, the workbook is pre-populated with the Safexpress rate
    card supplied by the ops team, ready to use as-is or edit for another
    courier / lane.
    """
    matrix_df = pd.DataFrame(
        _MATRIX_VALUES if prefill else [[None] * len(_ZONES) for _ in _ZONES],
        index=_ZONES, columns=_ZONES,
    )
    matrix_df.index.name = "From \\ To"

    mapping_rows = [
        {"Zone": z, "States / Locations (comma separated)": _ZONE_LOCATIONS[z] if prefill else ""}
        for z in _ZONES
    ]
    mapping_df = pd.DataFrame(mapping_rows)

    charge_rows = [
        {"Parameter": k, "Value": (v if prefill else "")}
        for k, v in _DEFAULT_CHARGES.items()
    ]
    charge_rows.append({
        "Parameter": "Metro Locations (comma separated)",
        "Value": _METRO_LOCATIONS if prefill else "",
    })
    charges_df = pd.DataFrame(charge_rows)

    out = io.BytesIO()
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        matrix_df.to_excel(writer, sheet_name="ZoneMatrix")
        mapping_df.to_excel(writer, sheet_name="ZoneMapping", index=False)
        charges_df.to_excel(writer, sheet_name="Charges", index=False)

        notes = pd.DataFrame({
            "Instructions": [
                "ZoneMatrix: per-kg rate from the zone in the row to the zone in the column.",
                "ZoneMapping: list every state / location that falls in each zone.",
                "Charges: flat/percentage charge parameters used by the price calculation.",
                "Appointment Charge Amount: flat rupee amount added when Appointment based",
                "  Delivery is selected for an AWB.",
                "Replace the sample Safexpress figures with your own agreement, or use as-is.",
                "Do not rename the sheet tabs (ZoneMatrix / ZoneMapping / Charges) or the tool cannot read them.",
            ]
        })
        notes.to_excel(writer, sheet_name="Read Me", index=False)

    return out.getvalue()


# --------------------------------------------------------------------------
# Price calculation
# --------------------------------------------------------------------------

def _num(x) -> Optional[float]:
    try:
        if x is None or (isinstance(x, float) and pd.isna(x)):
            return None
        return float(x)
    except (ValueError, TypeError):
        return None


def calculate_expected_price(
    pickup_state,
    drop_state,
    drop_city,
    ideal_weight,
    delivery_type: str,
    invoice_value,
    agreement: Agreement,
) -> dict:
    """Pure calculation of the expected (agreement-side) price for one AWB.

    Formula: Freight = max(Expected Weight × Agreement Rate/KG, Min
    Chargeable Freight); FSC = FSC% × Freight; FOV = max(FOV% × Invoice
    Value, FOV Minimum); Docket = flat; Metro = flat if the drop city is a
    metro location, else 0; Appointment = flat if Appointment based Delivery
    is selected, else 0. Calculated Total is the sum of all of these.

    Returns a dict of every line item. If a required input isn't available
    yet (no zone match, no ideal weight, no invoice value), `error` is set
    and `calculated_total` is None — the caller should treat that AWB as
    not yet calculable rather than guessing.
    """
    result = {
        "pickup_zone": None, "drop_zone": None, "agreement_rate": None,
        "expected_weight": None,
        "freight": None, "raw_freight": None, "freight_floor_applied": False,
        "fsc": None,
        "fov": None, "raw_fov": None, "fov_floor_applied": False,
        "docket": None,
        "metro_applied": False, "metro": 0.0,
        "appointment": 0.0,
        "calculated_total": None,
        "error": None,
    }

    pz = agreement.zone_for(pickup_state)
    dz = agreement.zone_for(drop_state)
    result["pickup_zone"] = pz
    result["drop_zone"] = dz
    if not pz or not dz:
        result["error"] = (
            "Could not derive a zone for the pickup/drop state — check the "
            "Commercial Agreement tab's Zone Mapping."
        )
        return result

    rate = agreement.rate_for(pz, dz)
    result["agreement_rate"] = rate
    if rate is None:
        result["error"] = f"No rate defined from {pz} to {dz} in the active agreement's Zone Matrix."
        return result

    ideal_weight_n = _num(ideal_weight)
    if ideal_weight_n is None or ideal_weight_n <= 0:
        result["error"] = "Enter the ideal (actual) weight first."
        return result

    invoice_value_n = _num(invoice_value)
    if invoice_value_n is None:
        result["error"] = "Invoice Value is missing for this AWB — needed to compute FOV."
        return result

    min_weight = agreement.charges.get("Min Chargeable Weight (Kg)") or 0
    expected_weight = max(ideal_weight_n, min_weight)
    result["expected_weight"] = expected_weight

    min_freight = agreement.charges.get("Min Chargeable Freight (Rs)") or 0
    raw_freight = expected_weight * rate
    freight = max(raw_freight, min_freight)
    result["raw_freight"] = raw_freight
    result["freight"] = freight
    result["freight_floor_applied"] = freight > raw_freight + 1e-9

    fsc_pct = agreement.charges.get("FSC Percent (%)") or 0
    fsc = freight * fsc_pct / 100.0
    result["fsc"] = fsc

    fov_pct = agreement.charges.get("FOV Percent (%)") or 0
    fov_min = agreement.charges.get("FOV Minimum (Rs)") or 0
    raw_fov = invoice_value_n * fov_pct / 100.0
    fov = max(raw_fov, fov_min)
    result["raw_fov"] = raw_fov
    result["fov"] = fov
    result["fov_floor_applied"] = fov > raw_fov + 1e-9

    docket = agreement.charges.get("Docket Charge (Flat Rs)") or 0
    result["docket"] = docket

    is_metro = agreement.is_metro(drop_city)
    metro_charge = agreement.charges.get("Metro Congestion Charge (Rs)") or 0
    metro = metro_charge if is_metro else 0.0
    result["metro_applied"] = is_metro
    result["metro"] = metro

    appt_amt = agreement.charges.get("Appointment Charge Amount (Rs)") or 0
    appointment = appt_amt if delivery_type == APPOINTMENT else 0.0
    result["appointment"] = appointment

    result["calculated_total"] = freight + fsc + fov + docket + metro + appointment
    return result


def billed_subtotal(
    freight_billed, fsc_billed, fov_billed, docket_billed,
    state_charges_billed, appointment_billed,
) -> Optional[float]:
    """Sum of the working sheet's billed components matching
    calculate_expected_price()'s line items exactly (Freight, FSC, FOV,
    Docket, State Charges [the Metro Congestion counterpart], Appointment).

    Deliberately excludes DHP, War Surcharge, ODA, Handling, Green Tax,
    Demurrage, and Other Charges — those aren't governed by the commercial
    agreement, so including them would make nearly every AWB look disputed
    regardless of whether the agreement-covered charges are correct.

    Returns None if any of the six components is missing (can't compare).
    """
    vals = [
        _num(x) for x in (
            freight_billed, fsc_billed, fov_billed, docket_billed,
            state_charges_billed, appointment_billed,
        )
    ]
    if any(v is None for v in vals):
        return None
    return sum(vals)


def decide_status(
    final_price: Optional[float],
    billed_total: Optional[float],
    tolerance: float = 1.0,
) -> tuple[str, str]:
    """Compare a resolved (approved-or-user-corrected) expected price
    against the billed subtotal. Returns (status, reason)."""
    if final_price is None:
        return STATUS_NOT_CHECKED, "Not yet reviewed."
    if billed_total is None:
        return STATUS_NOT_CHECKED, "Billed component columns missing — cannot compare."

    diff = final_price - billed_total
    if abs(diff) <= tolerance:
        return STATUS_APPROVED, ""

    direction = "over-billed by courier" if diff < 0 else "under-billed by courier"
    reason = (
        f"Expected ₹{final_price:,.2f} vs billed ₹{billed_total:,.2f} "
        f"(diff ₹{abs(diff):,.2f}, {direction})"
    )
    return STATUS_DISPUTED, reason


# --------------------------------------------------------------------------
# Output
# --------------------------------------------------------------------------

def to_download_bytes(df: pd.DataFrame, fmt: str) -> bytes:
    if fmt == "csv":
        return df.to_csv(index=False).encode("utf-8")

    out = io.BytesIO()
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="All AWBs", index=False)
        if "Status" in df.columns:
            df[df["Status"] == STATUS_APPROVED].to_excel(writer, sheet_name="Approved", index=False)
            df[df["Status"] == STATUS_DISPUTED].to_excel(writer, sheet_name="Disputed", index=False)

        # Highlight Disputed rows on the combined "All AWBs" sheet. Approved
        # rows are left with no special formatting, per spec.
        if "Status" in df.columns:
            ws = writer.sheets["All AWBs"]
            status_col_idx = list(df.columns).index("Status") + 1  # openpyxl is 1-indexed
            fill = PatternFill(start_color=DISPUTED_FILL, end_color=DISPUTED_FILL, fill_type="solid")
            for row_idx in range(2, ws.max_row + 1):  # row 1 is the header
                if ws.cell(row=row_idx, column=status_col_idx).value == STATUS_DISPUTED:
                    for col_idx in range(1, ws.max_column + 1):
                        ws.cell(row=row_idx, column=col_idx).fill = fill

    return out.getvalue()
